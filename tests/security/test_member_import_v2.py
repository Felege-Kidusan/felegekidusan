"""Member import v2 (2026-10-21): Excel Sync importer, templates, D5 Family Link rule.

Covers the owner-approved importer patch (plan rev 2):
  * all-or-nothing: preflight over the whole file, ONE transaction for
    members + codes + enrollments, dry_run ("Check file") writes nothing;
  * Ethiopian birth dates validated by round trip; date_of_birth + dob_ec_*;
  * age → age group → section for every new member (under 7 → ህጻናት);
  * strict class labels (never by numeric id; spaces ignored; ambiguity = error);
  * permanent → direct + code by age group; temporary → waiting list;
  * guardian mapping, emergency columns, System ID round trip, no duplicates;
  * GuardianService D5 sibling rule; profile edit keeps date_of_birth in step.

Static + PHP-CLI checks always run; live checks need the test server
(SSMS_TEST_BASE, default http://127.0.0.1:8080), sudo mariadb and openpyxl.
"""
import http.cookiejar
import io
import json
import os
import re
import shutil
import subprocess
import unittest
import urllib.error
import urllib.parse
import urllib.request
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
BASE = os.environ.get("SSMS_TEST_BASE", "http://127.0.0.1:8080").rstrip("/")
PASSWORD = os.environ.get("SSMS_TEST_PASSWORD", "AuditTest#2026")
HR_USER = os.environ.get("SSMS_TEST_HR_USER", "audit_hr")
PHP = shutil.which("php")


def read(rel):
    return (ROOT / rel).read_text(encoding="utf-8")


def php_eval(code):
    r = subprocess.run([PHP, "-r", code], capture_output=True, text=True, cwd=str(ROOT))
    if r.returncode != 0:
        raise AssertionError(r.stderr or r.stdout)
    return json.loads(r.stdout.strip().splitlines()[-1])


# ──────────────────────────────────────────────────────────────────────
# Static contracts
# ──────────────────────────────────────────────────────────────────────
class ImporterStaticContract(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.src = read("admin/api_import_members.php")
        cls.js = read("admin/js/hr_data_sync.js")
        cls.hr = read("admin/dashboards/hr-dept.php")
        cls.enroll = read("admin/backend/services/EnrollmentService.php")
        cls.map = read("admin/backend/services/ExcelColumnMap.php")

    def test_upload_hardening_kept(self):
        for token in ("is_uploaded_file(", "filesize(", "new ZipArchive()",
                      "$expandedBytes > 50 * 1024 * 1024", "$zip->numFiles > 5000",
                      "validateCsrf", "'super_admin', 'school_admin', 'hr_dept'"):
            self.assertIn(token, self.src)

    def test_single_transaction_and_dry_run(self):
        self.assertIn("$dryRun", self.src)
        self.assertEqual(self.src.count("$conn->begin_transaction()"), 1)
        self.assertIn("$conn->rollback()", self.src)
        self.assertIn("EnrollmentService::enroll($conn, $memberId, (int)$item['class']['id'], null, $adminId, true)", self.src)
        # The old post-commit enrollment pass and PDO writes are gone.
        self.assertNotIn("enrollByLabel(", self.src)
        self.assertNotIn("$pdo->beginTransaction()", self.src)
        # Writes happen only after the error gate and the dry-run exit.
        self.assertLess(self.src.index("if ($errorCount > 0)"), self.src.index("$conn->begin_transaction()"))
        self.assertLess(self.src.index("if ($dryRun)"), self.src.index("$conn->begin_transaction()"))

    def test_identity_and_section_rules(self):
        self.assertIn("$rowData['age_group']", self.src)
        self.assertNotIn("generateMemberCode($conn);", self.src)
        self.assertIn("normalizeSectionAm", self.src)
        self.assertIn("current_section", self.src)
        self.assertIn("MemberCategory::groupForAge($age)", self.src)
        self.assertIn("$rowData['registration_type'] = 'waiting'", self.src)
        self.assertIn("$rowData['registration_type'] = 'direct'", self.src)

    def test_strict_class_lookup(self):
        self.assertIn("EnrollmentService::resolveClassByLabel(", self.src)
        body = self.enroll[self.enroll.index("public static function resolveClassByLabel"):self.enroll.index("public static function activeClassesForLookup")]
        self.assertNotIn("ctype_digit", body)
        self.assertNotIn("WHERE id", body)
        self.assertIn("'ambiguous'", body)

    def test_ui_never_renders_server_text_as_html(self):
        self.assertIn("dry_run", self.js)
        self.assertIn("textContent", self.js)
        self.assertNotIn("${result.message}", self.js)
        self.assertNotIn("${file.name}</", self.js)
        for btn in ("btnCheckTemp", "btnCheckPerm"):
            self.assertIn(f'id="{btn}"', self.hr)
            self.assertIn(btn, self.js)
        self.assertIn("/admin/js/hr_data_sync.js?v=", self.hr)

    def test_templates(self):
        php = PHP or self.skipTest("PHP CLI is not installed")
        cols = php_eval(
            "require 'admin/backend/services/ExcelColumnMap.php';"
            "use App\\Services\\ExcelColumnMap as M;"
            "echo json_encode(['t'=>M::columns('temporary'),'p'=>M::columns('permanent'),"
            "'lt'=>M::locked('temporary'),'lp'=>M::locked('permanent'),"
            "'e1'=>M::resolveHeader('Emergency Contact'),'e2'=>M::resolveHeader('Emergency Phone'),"
            "'id'=>M::resolveHeader('ID'),'sid'=>M::resolveHeader('System ID')]);"
        )
        for key in ("class", "gender", "date_of_birth", "guardian_phone2", "emergency_name", "emergency_phone", "system_id"):
            self.assertIn(key, cols["t"])
        for key in ("emergency_name", "emergency_phone", "system_id", "member_code", "class"):
            self.assertIn(key, cols["p"])
        self.assertEqual(cols["lt"], ["system_id"])
        self.assertEqual(cols["lp"], ["member_code", "system_id"])
        self.assertEqual((cols["e1"], cols["e2"], cols["sid"]), ("emergency_name", "emergency_phone", "system_id"))
        self.assertIsNone(cols["id"], "a bare 'ID' column must not be read as the System ID")


# ──────────────────────────────────────────────────────────────────────
# PHP unit checks (pure helpers)
# ──────────────────────────────────────────────────────────────────────
@unittest.skipIf(PHP is None, "PHP CLI is not installed")
class ImportHelpersUnit(unittest.TestCase):
    def test_age_groups(self):
        out = php_eval(
            "require 'admin/backend/services/MemberCategory.php';"
            "use App\\Services\\MemberCategory as C; $o=[];"
            "foreach([0,5,6,7,13,14,17,18,44,151,-1] as $a){$o[(string)$a]=C::groupForAge($a);}"
            "$o['null']=C::groupForAge(null); echo json_encode($o);"
        )
        self.assertEqual(out, {"0": "7_13", "5": "7_13", "6": "7_13", "7": "7_13", "13": "7_13", "14": "14_17",
                               "17": "14_17", "18": "18_plus", "44": "18_plus", "151": None, "-1": None, "null": None})

    def test_ec_dates(self):
        out = php_eval(
            "require 'admin/backend/services/MemberImportService.php';"
            "use App\\Services\\MemberImportService as S;"
            "$t=new DateTimeImmutable('2026-10-21', new DateTimeZone('Africa/Addis_Ababa')); $o=[];"
            "foreach(['28/10/2006','2006-10-28','2006/10/28','04/13/2005','2011-13-06','2012-13-06','2010-13-06',"
            "'31/01/2000','2019-02-11','2019-02-12','10/28/2006','1899-01-01','abc',''] as $v){"
            "$r=S::parseEcDate($v,$t); $o[$v]=$r['ok']?$r['gregorian']:$r['error'];}"
            "echo json_encode($o);"
        )
        self.assertEqual(out["28/10/2006"], "2014-07-05")
        self.assertEqual(out["2006-10-28"], "2014-07-05")
        self.assertEqual(out["2006/10/28"], "2014-07-05")
        self.assertEqual(out["04/13/2005"], "2013-09-09")
        self.assertEqual(out["2011-13-06"], "2019-09-11")  # Pagume 6 in an EC leap year
        self.assertEqual(out["2012-13-06"], "day")
        self.assertEqual(out["2010-13-06"], "day")
        self.assertEqual(out["31/01/2000"], "day")
        self.assertEqual(out["2019-02-11"], "2026-10-21")  # today is allowed
        self.assertEqual(out["2019-02-12"], "future")
        self.assertEqual(out["10/28/2006"], "month")       # US order is refused, never guessed
        self.assertEqual(out["1899-01-01"], "year")
        self.assertEqual(out["abc"], "format")
        self.assertEqual(out[""], "empty")

    def test_cells_gender_names(self):
        out = php_eval(
            "require 'admin/backend/services/MemberImportService.php';"
            "use App\\Services\\MemberImportService as S;"
            "echo json_encode(['p'=>S::cleanCell(\"'+251911000000\"),'q'=>S::cleanCell(\"'abc\"),'t'=>S::cleanCell('  x '),'d'=>[S::cleanCell('-'),S::cleanCell(' -- '),S::cleanCell('—'),S::cleanCell(\"'-\"),S::cleanCell('-5')],"
            "'n'=>S::cleanName(\"  ሀ   ለ\\tሐ \"),'g'=>[S::normalizeGender('ወንድ'),S::normalizeGender('ሴት'),S::normalizeGender('M'),"
            "S::normalizeGender('female'),S::normalizeGender('x')],'s'=>S::splitName('ሀ ለ ሐ መ')]);"
        )
        self.assertEqual(out["p"], "+251911000000")  # export formula guard undone
        self.assertEqual(out["q"], "'abc")
        self.assertEqual(out["t"], "x")
        self.assertEqual(out["d"], ["", "", "", "", "-5"])  # dash placeholders = empty
        self.assertEqual(out["n"], "ሀ ለ ሐ")
        self.assertEqual(out["g"], ["male", "female", "male", "female", None])
        self.assertEqual(out["s"], ["ሀ", "ለ", "ሐ መ"])

    def test_class_label_resolution(self):
        out = php_eval(
            "require 'admin/backend/services/EnrollmentService.php';"
            "use App\\Services\\EnrollmentService as E;"
            "$c=[['id'=>5,'class_name'=>'1ኛ ክፍል','class_name_en'=>'Grade 1','class_code'=>'g1'],"
            "['id'=>9,'class_name'=>'ሣልሳይ (9ኛ ክፍል )','class_name_en'=>'','class_code'=>'s3'],"
            "['id'=>11,'class_name'=>'ቀዳማይ(7ኛ ክፍል)','class_name_en'=>'','class_code'=>'q1'],"
            "['id'=>12,'class_name'=>'7ኛ ክፍል','class_name_en'=>'','class_code'=>'g7'],"
            "['id'=>20,'class_name'=>'Dup','class_name_en'=>'','class_code'=>'d1'],"
            "['id'=>21,'class_name'=>'dup ','class_name_en'=>'','class_code'=>'d2']]; $o=[];"
            "foreach(['1ኛ ክፍል','1ኛክፍል','grade 1','G1','ሣልሳይ (9ኛ ክፍል)','ሣልሳይ(9ኛ ክፍል )','ቀዳማይ (7ኛ ክፍል)','7ኛ ክፍል','5','9','Dup','none'] as $l){"
            "$r=E::resolveClassByLabel(null,$l,$c); $o[$l]=$r['status']==='ok'?(int)$r['class']['id']:$r['status'];}"
            "echo json_encode($o);"
        )
        self.assertEqual(out["1ኛ ክፍል"], 5)
        self.assertEqual(out["1ኛክፍል"], 5)
        self.assertEqual(out["grade 1"], 5)
        self.assertEqual(out["G1"], 5)
        self.assertEqual(out["ሣልሳይ (9ኛ ክፍል)"], 9)
        self.assertEqual(out["ሣልሳይ(9ኛ ክፍል )"], 9)
        self.assertEqual(out["ቀዳማይ (7ኛ ክፍል)"], 11)
        self.assertEqual(out["7ኛ ክፍል"], 12)          # grade 7 and ቀዳማይ stay different classes
        self.assertEqual(out["5"], "unknown")          # never the class whose id is 5
        self.assertEqual(out["9"], "unknown")
        self.assertEqual(out["Dup"], "ambiguous")
        self.assertEqual(out["none"], "unknown")


# ──────────────────────────────────────────────────────────────────────
# Live
# ──────────────────────────────────────────────────────────────────────
def _server_up():
    try:
        urllib.request.urlopen(BASE + "/admin/index.php", timeout=3).read()
        return True
    except Exception:
        return False


def db(sql):
    r = subprocess.run(["sudo", "-n", "mariadb", "--default-character-set=utf8mb4", "-N", "-B", "ssms_test", "-e", sql],
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(r.stderr)
    return r.stdout.rstrip("\n")  # keep trailing empty (tab-separated) columns


try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    HAVE_OPENPYXL = False


class HrSession:
    def __init__(self):
        self.jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.jar))
        self.opener.addheaders = [("User-Agent", "import-v2-tests/1.0")]
        body = self._open(BASE + "/admin/index.php")[1]
        m = re.search(rb'name="csrf_token"\s+value="([a-f0-9]+)"', body)
        assert m, "csrf token not found on login page"
        data = urllib.parse.urlencode({"username": HR_USER, "password": PASSWORD, "csrf_token": m.group(1).decode()}).encode()
        status, body = self._open(BASE + "/admin/backend/login.php", data=data)
        if status == 429 or b"Too many login attempts" in body:
            raise unittest.SkipTest("admin login rate-limited — rerun later")
        page = self._open(BASE + "/admin/dashboards/hr-dept.php")[1].decode("utf-8", "replace")
        m2 = re.search(r"CSRF_TOKEN\s*=\s*'([0-9a-f]{64})'", page) or re.search(r'csrf_token" value="([0-9a-f]{64})', page)
        if not m2:
            raise unittest.SkipTest("HR test user cannot open the HR dashboard")
        self.csrf = m2.group(1)

    def _open(self, url, data=None, headers=None):
        req = urllib.request.Request(url, data=data)
        for k, v in (headers or {}).items():
            req.add_header(k, v)
        try:
            with self.opener.open(req, timeout=60) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()

    def upload(self, xlsx, tier="permanent", dry=False):
        boundary = "----imp" + uuid.uuid4().hex
        parts = []
        fields = [("csrf_token", self.csrf), ("tier", tier), ("dry_run", "1" if dry else "0")]
        if not dry:
            fields.append(("confirm_phrase", "IMPORT MEMBERS"))  # typed confirmation (danger-confirm DC-3)
        for k, v in fields:
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
        parts.append(
            f'--{boundary}\r\nContent-Disposition: form-data; name="import_file"; filename="m.xlsx"\r\n'
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n".encode()
            + xlsx + b"\r\n"
        )
        parts.append(f"--{boundary}--\r\n".encode())
        status, body = self._open(BASE + "/admin/api_import_members.php", data=b"".join(parts),
                                  headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        return status, json.loads(body.decode("utf-8"))

    def get(self, path, params=None):
        return self._open(BASE + path + ("?" + urllib.parse.urlencode(params) if params else ""))

    def post(self, path, fields):
        return self._open(BASE + path, data=urllib.parse.urlencode(fields).encode())


def book(headers, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["ፈለገ ቅዱሳን ሰንበት ትምህርት ቤት"])
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl is not installed")
class ImportLive(unittest.TestCase):
    TAG = "ቲቪ"  # every test member's grandfather name starts with this tag → easy cleanup

    @classmethod
    def setUpClass(cls):
        if not _server_up():
            raise unittest.SkipTest(f"test server not reachable at {BASE}")
        try:
            db("SELECT 1")
        except Exception as e:  # pragma: no cover
            raise unittest.SkipTest(f"sudo mariadb unavailable: {e}")
        db("DELETE FROM security_rate_limits") if db("SHOW TABLES LIKE 'security_rate_limits'") else None
        subprocess.run("rm -f /tmp/ssms_ratelimit/* 2>/dev/null", shell=True)
        cls.cleanup()
        # Two deliberately messy class names (irregular spaces), like production.
        db("INSERT INTO classes (class_name, class_name_en, class_code, is_active) VALUES "
           "('ሣልሳይ (9ኛ ክፍል )', '', 'tv_s3', 1), ('ቀዳማይ(7ኛ ክፍል)', '', 'tv_q1', 1)")
        cls.s = HrSession()

    @classmethod
    def tearDownClass(cls):
        cls.cleanup()

    @classmethod
    def cleanup(cls):
        ids = db(f"SELECT GROUP_CONCAT(id) FROM members WHERE grandfather_name LIKE '{cls.TAG}%'")
        if ids and ids != "NULL":
            db(f"DELETE FROM member_guardians WHERE member_id IN ({ids})")
            for table in ("class_enrollments", "member_changes_log", "member_history"):
                if db(f"SHOW TABLES LIKE '{table}'"):
                    db(f"DELETE FROM {table} WHERE member_id IN ({ids})")
            db(f"DELETE FROM members WHERE id IN ({ids})")
        db("DELETE g FROM guardians g LEFT JOIN member_guardians mg ON mg.guardian_id = g.id "
           f"WHERE mg.guardian_id IS NULL AND g.display_name LIKE '{cls.TAG}%'")
        db("DELETE FROM classes WHERE class_code IN ('tv_s3','tv_q1')")

    def member(self, name):
        row = db("SELECT id, IFNULL(member_code,''), gender, IFNULL(date_of_birth,''), "
                 "CONCAT_WS('-', dob_ec_year, dob_ec_month, dob_ec_day), IFNULL(age,''), IFNULL(age_group,''), "
                 "IFNULL(current_section,''), registration_type, membership_tier, IFNULL(waiting_since,''), "
                 "IFNULL(guardian_name,''), IFNULL(guardian_phone1,''), IFNULL(guardian_phone2,''), IFNULL(phone_guardian,''), "
                 "IFNULL(emergency_name,''), IFNULL(emergency_phone,''), IFNULL(phone_primary,''), "
                 "(SELECT IFNULL(GROUP_CONCAT(c.class_code),'') FROM class_enrollments ce JOIN classes c ON c.id = ce.class_id "
                 " WHERE ce.member_id = m.id AND ce.status = 'active') "
                 f"FROM members m WHERE full_name_am = '{name}'")
        self.assertTrue(row, f"member {name} not found")
        keys = ["id", "code", "gender", "dob", "ec", "age", "group", "section", "reg", "tier", "waiting",
                "gname", "gp1", "gp2", "pg", "ename", "ephone", "phone", "classes"]
        self.assertEqual(len(row.splitlines()), 1, f"duplicate rows for {name}")
        return dict(zip(keys, row.split("\t")))

    def count(self):
        return int(db(f"SELECT COUNT(*) FROM members WHERE grandfather_name LIKE '{self.TAG}%'"))

    H = ["Full Name", "Christian Name", "Class", "Gender", "Date of Birth", "Primary Phone", "Guardian Name",
         "Guardian Phone 1", "Guardian Phone 2", "Emergency Contact", "Emergency Phone"]

    def test_1_bad_file_is_rejected_whole(self):
        before = self.count()
        f = book(self.H, [
            ["ሀሌታ አበራ ቲቪአንድ", "", "ሣልሳይ (9ኛ ክፍል)", "ሴት", "28/10/2006", "", "", "", "", "", ""],
            ["ለማ አበራ ቲቪሁለት", "", "3", "ወንድ", "2010-13-06", "", "", "", "", "", ""],
            ["ሐና አበራ ቲቪሶስት", "", "", "X", "05/01/2012", "", "", "", "", "", ""],
            ["ሀሌታ አበራ ቲቪአንድ", "", "", "ሴት", "28/10/2006", "", "", "", "", "", ""],
        ])
        for dry in (True, False):
            status, res = self.s.upload(f, dry=dry)
            self.assertEqual(status, 200)
            self.assertFalse(res["success"])
            joined = "\n".join(res["errors"])
            self.assertIn("Row 5:", joined)  # Pagume 6 in a non-leap year + class "3"
            self.assertIn('Class "3" is not an active class', joined)
            self.assertIn("Row 6: Gender", joined)
            self.assertIn("Row 7: same member as row 4", joined)
        status, res = self.s.upload(book(["Full Name", "Gender", "Status"], [["አርካይቭ ሙከራ ቲቪአር", "ወንድ", "Archived"]]), dry=True)
        self.assertFalse(res["success"])
        self.assertIn("cannot be set from Excel", res["errors"][0])
        self.assertEqual(self.count(), before, "a rejected file must save nothing")

    def test_2_import_check_then_save_then_reupload(self):
        f = book(self.H, [
            ["ሀሌታ አበራ ቲቪአንድ", "ወለተ ማርያም", "ሣልሳይ(9ኛ ክፍል )", "ሴት", "28/10/2006", "",
             "ቲቪ እናት", "+251911000001", "0911000002", "ጎረቤት", "0911000003"],
            ["ለማ አበራ ቲቪሁለት", "", " ቀዳማይ (7ኛ ክፍል) ", "M", "04/13/2003", "'+251911000009", "", "", "", "", ""],
            ["ሐና አበራ ቲቪሶስት", "", "", "F", "01/01/2013", "", "", "", "", "", ""],
            ["ዮናስ አበራ ቲቪአራት", "", "", "ወንድ", "01/01/1990", "", "", "", "", "", ""],
        ])
        before = self.count()
        status, res = self.s.upload(f, dry=True)
        self.assertTrue(res["success"], res)
        self.assertTrue(res["dry_run"])
        self.assertEqual(res["summary"]["new_permanent"], 4)
        self.assertEqual(res["summary"]["to_enroll"], 2)
        self.assertEqual(self.count(), before, "Check file must save nothing")

        status, res = self.s.upload(f)
        self.assertTrue(res["success"], res)
        self.assertEqual(res["stats"]["inserted"], 4)
        self.assertEqual(res["stats"]["enrolled"], 2)
        self.assertEqual(res["stats"]["unassigned"], 2)

        a = self.member("ሀሌታ አበራ ቲቪአንድ")
        self.assertEqual((a["gender"], a["dob"], a["ec"]), ("female", "2014-07-05", "2006-10-28"))
        self.assertEqual((a["age"], a["group"], a["section"]), ("12", "7_13", "ህጻናት"))
        self.assertEqual((a["reg"], a["tier"]), ("direct", "permanent"))
        self.assertTrue(a["code"].startswith("A-"), a["code"])
        self.assertEqual((a["gname"], a["gp1"], a["gp2"], a["pg"]), ("ቲቪ እናት", "+251911000001", "0911000002", "+251911000001"))
        self.assertEqual((a["ename"], a["ephone"]), ("ጎረቤት", "0911000003"))
        self.assertEqual(a["classes"], "tv_s3")

        b = self.member("ለማ አበራ ቲቪሁለት")
        self.assertEqual(b["phone"], "+251911000009")  # export guard apostrophe removed
        self.assertEqual(b["ec"], "2003-13-4")
        self.assertEqual(b["group"], "14_17")
        self.assertTrue(b["code"].startswith("B-"), b["code"])
        self.assertEqual(b["classes"], "tv_q1")

        c = self.member("ሐና አበራ ቲቪሶስት")  # age 6 → ህጻናት (owner D3)
        self.assertEqual((c["age"], c["group"], c["section"], c["classes"]), ("6", "7_13", "ህጻናት", ""))
        d = self.member("ዮናስ አበራ ቲቪአራት")
        self.assertEqual((d["group"], d["section"]), ("18_plus", "ወጣቶች"))
        self.assertTrue(d["code"].startswith("C-"), d["code"])

        status, res = self.s.upload(f)
        self.assertTrue(res["success"], res)
        self.assertEqual(res["stats"]["inserted"], 0)
        self.assertEqual(res["stats"]["updated"], 0)
        self.assertEqual(self.count(), before + 4)

    def test_3_temporary_waiting_list_and_class(self):
        f = book(["Full Name", "Class", "Gender", "Date of Birth", "Guardian Name", "Guardian Phone 1"], [
            ["ጊዜ አንድ ቲቪጊዜ", "ቀዳማይ(7ኛ ክፍል)", "ወንድ", "10/02/2010", "ቲቪ እማማ", "0922000000"],
        ])
        status, res = self.s.upload(f, tier="temporary")
        self.assertTrue(res["success"], res)
        t = self.member("ጊዜ አንድ ቲቪጊዜ")
        self.assertEqual((t["code"], t["reg"], t["tier"]), ("", "waiting", "temporary"))
        self.assertRegex(t["waiting"], r"^\d{4}-\d{2}-\d{2}$")
        self.assertEqual(t["classes"], "tv_q1")
        self.assertEqual(t["pg"], "0922000000")

    def test_4_exported_template_round_trip_and_duplicate_guard(self):
        from openpyxl import load_workbook
        status, body = self.s.get("/admin/api_export_members.php", {"tier": "temporary"})
        self.assertEqual(status, 200)
        ws = load_workbook(io.BytesIO(body)).active
        headers = [c.value for c in ws[3]]
        self.assertIn("System ID", headers)
        self.assertIn("Class", headers)
        status, res = self.s.upload(body, tier="temporary", dry=True)
        self.assertTrue(res["success"], res)
        self.assertEqual(res["summary"]["new_temporary"] + res["summary"]["new_permanent"], 0)
        # Same name again, no birth date → refused, never a silent duplicate.
        f = book(["Full Name", "Gender"], [["ጊዜ አንድ ቲቪጊዜ", "ወንድ"]])
        status, res = self.s.upload(f, tier="temporary", dry=True)
        self.assertFalse(res["success"])
        self.assertIn("already exists", res["errors"][0])

    def test_5_rollback_when_enrollment_fails(self):
        before = self.count()
        db("RENAME TABLE class_enrollments TO class_enrollments_tv_tmp")
        try:
            f = book(["Full Name", "Class", "Gender", "Date of Birth"], [
                ["ሮል አንድ ቲቪሮል", "", "ወንድ", "01/01/2005"],
                ["ሮል ሁለት ቲቪሮል", "ቀዳማይ(7ኛ ክፍል)", "ሴት", "01/01/2006"],
            ])
            status, res = self.s.upload(f)
        finally:
            db("RENAME TABLE class_enrollments_tv_tmp TO class_enrollments")
        self.assertFalse(res["success"], res)
        self.assertIn("Nothing was saved", res["message"])
        self.assertEqual(self.count(), before, "row 4 must be rolled back when row 5 fails")

    def test_6_d5_family_link_rule(self):
        f = book(["Full Name", "Gender", "Date of Birth", "Guardian Name"], [
            ["ወንድም ከበደ ቲቪቤተሰብ", "ወንድ", "01/01/2008", "ቲቪ ማርታ ተስፋዬ"],
            ["እህት ከበደ ቲቪቤተሰብ", "ሴት", "01/01/2010", "ቲቪ ማርታ  ተስፋዬ"],
            ["ሌላ ከበደ ቲቪሌላ", "ሴት", "01/01/2011", "ቲቪ ማርታ ተስፋዬ"],
        ])
        status, res = self.s.upload(f)
        self.assertTrue(res["success"], res)
        cursor = 0
        for _ in range(50):
            status, body = self.s.post("/admin/api_guardians.php", {"action": "attach_batch", "cursor": cursor, "confirm_phrase": "LINK ALL", "csrf_token": self.s.csrf})
            r = json.loads(body.decode("utf-8"))
            if r.get("done"):
                break
            cursor = r.get("next_cursor", 0)
        g = {name: db("SELECT IFNULL(MAX(mg.guardian_id),0) FROM members m LEFT JOIN member_guardians mg ON mg.member_id = m.id "
                      f"WHERE m.full_name_am = '{name}'")
             for name in ("ወንድም ከበደ ቲቪቤተሰብ", "እህት ከበደ ቲቪቤተሰብ", "ሌላ ከበደ ቲቪሌላ")}
        self.assertNotEqual(g["ወንድም ከበደ ቲቪቤተሰብ"], "0")
        self.assertEqual(g["ወንድም ከበደ ቲቪቤተሰብ"], g["እህት ከበደ ቲቪቤተሰብ"], "siblings must share the guardian (D5)")
        self.assertNotEqual(g["ሌላ ከበደ ቲቪሌላ"], g["ወንድም ከበደ ቲቪቤተሰብ"], "different grandfather → not auto-linked")

    def test_7_profile_edit_keeps_gregorian_birth_date_in_step(self):
        f = book(["Full Name", "Gender", "Date of Birth"], [["ፕሮፋይል አንድ ቲቪፕሮ", "ሴት", "01/01/2010"]])
        status, res = self.s.upload(f)
        self.assertTrue(res["success"], res)
        m = self.member("ፕሮፋይል አንድ ቲቪፕሮ")
        self.assertEqual(m["dob"], "2017-09-11")
        status, body = self.s.post("/admin/info_manage_member.php", {
            "csrf_token": self.s.csrf, "id": m["id"], "dob_day": "28", "dob_month": "10", "dob_year": "2006"})
        self.assertEqual(json.loads(body.decode("utf-8")).get("status"), "success", body[:300])
        m = self.member("ፕሮፋይል አንድ ቲቪፕሮ")
        self.assertEqual((m["ec"], m["dob"], m["age"]), ("2006-10-28", "2014-07-05", "12"))
        # An impossible EC date never overwrites the stored Gregorian date.
        self.s.post("/admin/info_manage_member.php", {
            "csrf_token": self.s.csrf, "id": m["id"], "dob_day": "6", "dob_month": "13", "dob_year": "2010"})
        self.assertEqual(self.member("ፕሮፋይል አንድ ቲቪፕሮ")["dob"], "2014-07-05")

    def test_8_link_all_walks_every_step(self):
        # Regression (2026-10-21): "Link all" used to stop after the first
        # step of 40 because the first step read the newest members.
        letters = "ሀለሐመሠረሰሸቀበተቸኀነኘአከኸወዐዘዠየደጀገጠጨጰጸፀፈፐ"
        rows = [[f"ልጅ{letters[i % len(letters)]}{i} አባት{i} ቲቪሊንክ{i}", "ወንድ", "01/01/2010", f"ቲቪ እናት {i}"] for i in range(45)]
        status, res = self.s.upload(book(["Full Name", "Gender", "Date of Birth", "Guardian Name"], rows))
        self.assertTrue(res["success"], res)
        cursor, steps = 0, 0
        while steps < 50:
            status, body = self.s.post("/admin/api_guardians.php", {"action": "attach_batch", "cursor": cursor, "confirm_phrase": "LINK ALL", "csrf_token": self.s.csrf})
            r = json.loads(body.decode("utf-8"))
            steps += 1
            if r.get("done") or int(r.get("processed", 0)) == 0:  # same stop rule as family_link.js
                break
            cursor = r.get("next_cursor", 0)
        unlinked = db("SELECT COUNT(*) FROM members m WHERE m.grandfather_name LIKE 'ቲቪሊንክ%' "
                      "AND NOT EXISTS (SELECT 1 FROM member_guardians mg WHERE mg.member_id = m.id)")
        self.assertEqual(unlinked, "0", f"Link all left members unlinked after {steps} step(s)")
        self.assertGreaterEqual(steps, 2)


if __name__ == "__main__":
    unittest.main()
